import base64
import glob
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timedelta
import numpy as np
import openpyxl
import pandas as pd
import plotly.graph_objects as go
import plotly.io as pio
import streamlit as st

TARGET_DATE = datetime.now().date() - timedelta(days=1)

# Full absolute paths for local plant environment
LOCAL_FILE_PATH = (
    r"C:\Users\AUB5367\OneDrive - MDLZ\DMS-2\PERK MIS.-2026.xlsx"
)
FILE_PATH = (
    LOCAL_FILE_PATH
    if os.path.exists(LOCAL_FILE_PATH)
    else "PERK MIS.-2026.xlsx"
)

SHEET_DMS = "DMS-2"
MINOR_STOPS_DIR = r"C:\Users\AUB5367\Desktop\New folder\AM STEP-4\dESKTOP BACKUP\IL6S\Minor Stop-SWP\MTBF-PBI 2026\Minor Stops"
SHOPLOGIX_DIR = r"c:\Users\AUB5367\OneDrive - MDLZ\Desktop\New folder\AM STEP-4\dESKTOP BACKUP\IL6S\Minor Stop-SWP\MTBF-PBI 2026\DMS-Downtimes"
EFFICIENCY_DIR = r"C:\Users\AUB5367\OneDrive - MDLZ\Desktop\New folder\AM STEP-4\dESKTOP BACKUP\IL6S\Minor Stop-SWP\MTBF-PBI 2026\Machine Efficiency"
DOWNLOAD_DIR = r"C:\Users\AUB5367\Downloads"
MOM_FILE_PATH = (
    r"C:\Users\AUB5367\OneDrive - MDLZ\Desktop\Python Automations\Daily MTBF\PERK_MOM_History.xlsx"
    if os.path.exists(
        r"C:\Users\AUB5367\OneDrive - MDLZ\Desktop\Python Automations\Daily MTBF\PERK_MOM_History.xlsx"
    )
    else "PERK_MOM_History.xlsx"
)

SHOPLOGIX_MAKING_URL = (
    "https://portal.shoplogix.com/app/main/dashboards/6a86963d09213c5b9401f708"
)
SHOPLOGIX_EFF_URL = (
    "https://portal.shoplogix.com/app/main/dashboards/66af6f6e1f22b100361259da"
)
SHOPLOGIX_MINOR_STOPS_URL = (
    "https://portal.shoplogix.com/app/main/dashboards/68580fd5651fa00033608377"
)


def safe_copy_file(src):
  if not os.path.exists(src):
    return src
  temp_dir = tempfile.gettempdir()
  dst = os.path.join(temp_dir, f"temp_{os.path.basename(src)}")
  try:
    subprocess.run(
        f'copy /Y "{src}" "{dst}"',
        shell=True,
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
  except Exception:
    try:
      shutil.copy2(src, dst)
    except Exception:
      return src
  return dst


def to_norm_tag(val):
  if pd.isna(val) or str(val).strip() in ("", "nan", "NAN"):
    return ""
  val_s = str(val).strip()
  dt = pd.to_datetime(val_s, errors="coerce")
  if pd.notna(dt):
    return f"{dt.day}-{dt.strftime('%b')}".upper()
  m = re.match(r"^(\d{1,2})[-/ ]*([A-Za-z]{3})", val_s)
  if m:
    return f"{int(m.group(1))}-{m.group(2)[:3]}".upper()
  return val_s.upper()


def parse_date_to_timestamp(val_s):
  if pd.isna(val_s) or not str(val_s).strip():
    return pd.Timestamp.min
  s = str(val_s).strip()
  dt = pd.to_datetime(s, errors="coerce")
  if pd.notna(dt):
    return dt
  try:
    return pd.to_datetime(s + "-2026", format="%d-%b-%Y")
  except Exception:
    return pd.Timestamp.min


def upload_file_to_github(file_path, repo_path):
  token = "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"
  repo = "aub5367/Perk-DMS-Board"
  url = f"https://api.github.com/repos/{repo}/contents/{repo_path}"

  if not os.path.exists(file_path):
    return

  with open(file_path, "rb") as f:
    content_bytes = f.read()

  encoded_content = base64.b64encode(content_bytes).decode("utf-8")
  headers = {
      "Authorization": f"Bearer {token}",
      "Accept": "application/vnd.github+json",
  }

  get_resp = requests.get(url, headers=headers)
  sha = get_resp.json().get("sha") if get_resp.status_code == 200 else None

  payload = {
      "message": f"Auto-update {repo_path} from plant PC",
      "content": encoded_content,
  }
  if sha:
    payload["sha"] = sha

  put_resp = requests.put(url, headers=headers, json=payload)
  if put_resp.status_code in [200, 201]:
    print(f"✅ Successfully uploaded {repo_path} to GitHub!")
  else:
    print(f"⚠️ Failed to upload {repo_path}: {put_resp.text}")


# SELF-LAUNCHING TRIGGER
is_streamlit_running = (
    "streamlit" in sys.modules
    or os.environ.get("STREAMLIT_SERVER_PORT") is not None
)

if not is_streamlit_running and os.name == "nt":
  print("=" * 60)
  print("Step 1: Connecting to Shoplogix Portal...")
  print("=" * 60)

  try:
    from selenium import webdriver
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.support.ui import WebDriverWait

    driver = webdriver.Edge()
    wait = WebDriverWait(driver, 45)

    def purge_folder(folder_path):
      if os.path.exists(folder_path):
        for f in os.listdir(folder_path):
          if f.endswith(".xlsx") or f.endswith(".xls"):
            try:
              os.remove(os.path.join(folder_path, f))
            except Exception:
              pass

    def download_report(dest_folder, report_prefix):
      os.makedirs(dest_folder, exist_ok=True)
      try:
        wait.until(
            EC.invisibility_of_element_located(
                (By.CSS_SELECTOR, ".widget-refresh-overlay.widget-loading")
            )
        )
      except Exception:
        pass

      time.sleep(3)
      print(f"📥 Exporting report for {report_prefix}...")
      edit_btn = wait.until(
          EC.presence_of_element_located((By.XPATH, "//button[contains(@title,'Edit')]"))
      )
      driver.execute_script("arguments[0].scrollIntoView(true);", edit_btn)
      driver.execute_script("arguments[0].click();", edit_btn)

      dl_btn = wait.until(
          EC.presence_of_element_located(
              (By.CSS_SELECTOR, "button.js--btn-download-menu")
          )
      )
      driver.execute_script("arguments[0].click();", dl_btn)

      excel_btn = wait.until(
          EC.presence_of_element_located(
              (By.XPATH, "//*[normalize-space()='Excel File']")
          )
      )
      driver.execute_script("arguments[0].click();", excel_btn)

      time.sleep(12)
      files = [
          os.path.join(DOWNLOAD_DIR, f)
          for f in os.listdir(DOWNLOAD_DIR)
          if f.endswith(".xlsx") and not f.startswith("~$")
      ]
      if files:
        latest = max(files, key=os.path.getctime)
        purge_folder(dest_folder)
        target_filename = (
            f"{report_prefix}_{TARGET_DATE.strftime('%Y_%m_%d')}.xlsx"
        )
        shutil.move(latest, os.path.join(dest_folder, target_filename))
        print(f"✅ Successfully downloaded and updated: {target_filename}")

    driver.get(SHOPLOGIX_MAKING_URL)
    try:
      user_input = wait.until(
          EC.visibility_of_element_located((By.NAME, "Username"))
      )
      user_input.send_keys("himanshu.chauhan@mdlz.com")
      driver.find_element(By.XPATH, "//button[text()='Next']").click()
      print(
          "ℹ️ Username entered. Please complete password/MFA in the browser"
          " window if prompted..."
      )
    except Exception:
      print("ℹ️ Already logged in or login page loaded differently. Proceeding...")

    time.sleep(15)
    download_report(SHOPLOGIX_DIR, "Shoplogix_Making_Downtimes")

    driver.get(SHOPLOGIX_EFF_URL)
    time.sleep(8)
    download_report(EFFICIENCY_DIR, "Machine_Efficiency")

    driver.get(SHOPLOGIX_MINOR_STOPS_URL)
    time.sleep(8)
    download_report(MINOR_STOPS_DIR, "Shoplogix_Minor_Stops")

    driver.quit()
  except Exception as e:
    print(f"⚠️ Automation download encountered an issue: {e}")

  print("📤 Syncing local files to GitHub via API...")
  upload_file_to_github(FILE_PATH, "PERK MIS.-2026.xlsx")
  if os.path.exists(MOM_FILE_PATH):
    upload_file_to_github(MOM_FILE_PATH, "PERK_MOM_History.xlsx")

  print("🚀 Launching Dashboard Viewer...")
  subprocess.run([sys.executable, "-m", "streamlit", "run", __file__])
  sys.exit(0)

st.set_page_config(
    page_title="PERK DMS BOARD", layout="wide", initial_sidebar_state="expanded"
)

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&display=swap');
    html, body, [class*="css"] {
        font-family: 'Inter', -apple-system, BlinkMacSystemFont, sans-serif;
    }
    .board-header {
        background: linear-gradient(135deg, #1f0b35 0%, #3b1764 50%, #54228a 100%);
        color: #ffffff;
        text-align: center;
        padding: 16px 24px;
        border-radius: 12px;
        font-weight: 800;
        font-size: 2rem;
        letter-spacing: 1.5px;
        box-shadow: 0 6px 18px rgba(31, 11, 53, 0.3);
        margin-bottom: 24px;
        border: 1px solid rgba(255, 255, 255, 0.12);
    }
    .dms-table-container {
        width: 100%;
        overflow-x: auto;
        border-radius: 12px;
        border: 2px solid #2a1147;
        box-shadow: 0 8px 24px rgba(0, 0, 0, 0.08);
        background: #ffffff;
        margin-bottom: 30px;
    }
    table.dms-table {
        width: 100%;
        border-collapse: separate;
        border-spacing: 0;
        font-size: 15px;
        text-align: center;
        background-color: #ffffff;
    }
    table.dms-table th, table.dms-table td {
        border-right: 1px solid #e2e8f0;
        border-bottom: 1px solid #e2e8f0;
        padding: 10px 12px;
    }
    table.dms-table thead th {
        background: #250e3e;
        color: #ffffff;
        font-weight: 700;
        font-size: 16px;
        text-transform: uppercase;
        letter-spacing: 0.5px;
        padding: 12px 10px;
    }
    .th-sub {
        background: #3c1964 !important;
        color: #f1f5f9 !important;
        font-size: 13px !important;
        font-weight: 700 !important;
        letter-spacing: 0.8px;
        padding: 7px 6px !important;
    }
    .cell-pillar {
        font-weight: 800;
        font-size: 15px;
        letter-spacing: 0.8px;
        vertical-align: middle;
        text-align: center;
        border-right: 2px solid #2a1147 !important;
    }
    .pillar-safety { background-color: #fffbeb !important; color: #b45309 !important; }
    .pillar-quality { background-color: #eff6ff !important; color: #1d4ed8 !important; }
    .pillar-il6s { background-color: #fefce8 !important; color: #854d0e !important; }
    .pillar-delivery { background-color: #f0fdf4 !important; color: #15803d !important; }
    .pillar-cost { background-color: #faf5ff !important; color: #6b21a8 !important; }
    .pillar-sustainability { background-color: #ecfdf5 !important; color: #047857 !important; }
    .pillar-morale { background-color: #fff1f2 !important; color: #be123c !important; }
    .cell-kpi {
        text-align: left !important;
        font-weight: 600;
        font-size: 15px;
        color: #1e293b;
        background-color: #ffffff;
        white-space: nowrap;
        padding-left: 16px !important;
    }
    .cell-uom {
        color: #64748b;
        font-weight: 600;
        font-size: 14px;
        background-color: #f8fafc;
    }
    .cell-target {
        background-color: #334155 !important;
        color: #ffffff !important;
        font-weight: 700;
        font-size: 15px;
    }
    .card-container {
        border: 2px solid #2a1147;
        border-radius: 10px;
        background-color: #ffffff;
        margin-bottom: 14px;
        box-shadow: 0 4px 14px rgba(0,0,0,0.06);
        text-align: center;
    }
    .card-header {
        background-color: #250e3e;
        color: #ffffff;
        font-weight: 700;
        font-size: 12px;
        padding: 7px;
        letter-spacing: 0.8px;
        border-top-left-radius: 8px;
        border-top-right-radius: 8px;
    }
    .card-body {
        padding: 8px;
        font-size: 18px;
        font-weight: 800;
    }
    .chart-card {
        background: #ffffff;
        border-radius: 12px;
        border: 1px solid #e2e8f0;
        padding: 20px;
        box-shadow: 0 4px 14px rgba(0,0,0,0.04);
    }
    </style>
""",
    unsafe_allow_html=True,
)

with st.sidebar:
  st.markdown("### ⚙️ Dashboard Controls")
  st.markdown("---")
  st.markdown("#### 📂 Manual File Uploaders")
  uploaded_mis = st.file_uploader(
      "Upload `PERK MIS.-2026.xlsx`", type=["xlsx"], key="mis_uploader"
  )
  uploaded_shoplogix = st.file_uploader(
      "Upload Shoplogix Making Report", type=["xlsx", "xls"], key="slx_uploader"
  )
  uploaded_eff = st.file_uploader(
      "Upload Machine Efficiency Report", type=["xlsx", "xls"], key="eff_uploader"
  )
  uploaded_minor_stops = st.file_uploader(
      "Upload Shoplogix Minor Stops Report",
      type=["xlsx", "xls"],
      key="ms_uploader",
  )
  uploaded_mom = st.file_uploader(
      "Upload `PERK_MOM_History.xlsx`", type=["xlsx"], key="mom_uploader"
  )


@st.cache_data
def load_dms_6day_matrix(path_or_buffer, sheet):
  if hasattr(path_or_buffer, "read"):
    wb = openpyxl.load_workbook(path_or_buffer, data_only=True)
    ws = wb[sheet]
    raw_list = list(ws.iter_rows(values_only=True))
    raw = pd.DataFrame(raw_list)
  else:
    if not os.path.exists(path_or_buffer):
      raise FileNotFoundError(f"File not found: {path_or_buffer}")
    tmp_path = safe_copy_file(path_or_buffer)
    try:
      wb = openpyxl.load_workbook(tmp_path, data_only=True)
      ws = wb[sheet]
      raw_list = list(ws.iter_rows(values_only=True))
      raw = pd.DataFrame(raw_list)
    finally:
      if os.path.exists(tmp_path) and tmp_path != path_or_buffer:
        try:
          os.remove(tmp_path)
        except Exception:
          pass

  try:
    if hasattr(path_or_buffer, "read"):
      wb_styled = openpyxl.load_workbook(path_or_buffer, data_only=True)
    else:
      wb_styled = openpyxl.load_workbook(tmp_path, data_only=True)
    ws_styled = wb_styled[sheet]
  except Exception:
    ws_styled = None

  target_row, target_start_col = 3, 3
  for r in range(min(10, raw.shape[0])):
    for c in range(raw.shape[1]):
      val = str(raw.iat[r, c]).strip().upper()
      if "TARGET" in val or "ACTUAL" in val:
        target_row, target_start_col = r, c
        break
    if target_row != 3 or target_start_col != 3:
      break

  date_row = max(0, target_row - 2)
  day_row = max(0, target_row - 1)
  data_start_row = target_row + 1

  date_blocks = []
  c = target_start_col
  while c < raw.shape[1] - 1:
    t_header = str(raw.iat[target_row, c]).strip().upper()
    if "TARGET" not in t_header and not pd.isna(raw.iat[target_row, c]):
      c += 1
      continue

    date_val = raw.iat[date_row, c]
    if pd.isna(date_val) or str(date_val).strip() in ("", "nan", "NAN"):
      if c > 0:
        date_val = raw.iat[date_row, c - 1]

    if pd.isna(date_val) or str(date_val).strip() in ("", "nan", "NAN"):
      c += 1
      continue

    date_str = str(date_val).split(" ")[0].strip()
    day_val = raw.iat[day_row, c]
    if pd.isna(day_val) or str(day_val).strip() in ("", "nan", "NAN"):
      if c > 0:
        day_val = raw.iat[day_row, c - 1]
    day_str = str(day_val).strip() if not pd.isna(day_val) else ""

    date_blocks.append({
        "target_col": c,
        "actual_col": c + 1,
        "date": date_str,
        "day": day_str,
        "norm_tag": to_norm_tag(date_val),
        "timestamp": parse_date_to_timestamp(date_str),
    })
    c += 2

  recent_blocks = date_blocks[-6:] if len(date_blocks) >= 6 else date_blocks

  KNOWN_PILLARS = [
      "SAFETY",
      "QUALITY",
      "IL6S",
      "DELIVERY",
      "COST",
      "SUSTAINABILITY",
      "MORALE",
  ]
  KNOWN_UOMS = ["NOS", "MT", "%", "YES/NO", "HRS", "MMBTU"]

  current_pillar = "SAFETY"
  rows_data = []

  for r in range(data_start_row, raw.shape[0]):
    left_vals = []
    for col_i in range(target_start_col):
      v = raw.iat[r, col_i]
      if not pd.isna(v) and str(v).strip() not in ("", "nan", "NAN"):
        left_vals.append(str(v).strip())

    if not left_vals:
      continue

    detected_pillar = None
    for val in left_vals:
      for p in KNOWN_PILLARS:
        if p in val.upper():
          detected_pillar = p
          break
      if detected_pillar:
        break

    if detected_pillar:
      current_pillar = detected_pillar
      left_vals = [v for v in left_vals if detected_pillar not in v.upper()]

    if not left_vals:
      continue

    uom = ""
    for val in left_vals:
      if val.upper() in KNOWN_UOMS:
        uom = val
        left_vals.remove(val)
        break

    kpi_name = left_vals[0] if left_vals else ""
    if not kpi_name:
      continue

    row_item = {"pillar": current_pillar, "kpi": kpi_name, "uom": uom, "days": []}

    for block in recent_blocks:
      t_col = block["target_col"]
      a_col = block["actual_col"]
      t_val = raw.iat[r, t_col] if t_col < raw.shape[1] else ""
      a_val = raw.iat[r, a_col] if a_col < raw.shape[1] else ""

      excel_bg = ""
      if ws_styled:
        try:
          cell = ws_styled.cell(row=r + 1, column=a_col + 1)
          fill = cell.fill
          if fill and fill.fill_type and fill.start_color:
            color_val = fill.start_color.rgb
            if color_val and isinstance(color_val, str) and len(color_val) == 8:
              excel_bg = f"#{color_val[2:]}"
            elif (
                color_val
                and isinstance(color_val, str)
                and len(color_val) == 6
            ):
              excel_bg = f"#{color_val}"
        except Exception:
          pass

      row_item["days"].append({
          "target": "" if pd.isna(t_val) else str(t_val).strip(),
          "actual": "" if pd.isna(a_val) else str(a_val).strip(),
          "excel_bg": excel_bg,
      })

    rows_data.append(row_item)

  return rows_data, recent_blocks, date_blocks


@st.cache_data
def load_ge_loss_monthly(path):
  return {
      "JAN": 81.5,
      "FEB": 80.5,
      "MAR": 80.6,
      "APR": 80.1,
      "MAY": 81.1,
      "JUN": 67.2,
      "JUL": 75.0,
      "AUG": 46.3,
      "SEP": "",
      "OCT": "",
      "NOV": "",
      "DEC": "",
  }


@st.cache_data
def load_shoplogix_from_folder(folder_path):
  if not os.path.exists(folder_path):
    os.makedirs(folder_path, exist_ok=True)
    return pd.DataFrame()
  files = glob.glob(os.path.join(folder_path, "**", "*.xls*"), recursive=True)
  if not files:
    return pd.DataFrame()
  latest_file = max(files, key=os.path.getmtime)
  tmp_path = safe_copy_file(latest_file)
  try:
    df = pd.read_excel(tmp_path)
    return standardize_slx_df(df)
  except Exception:
    return pd.DataFrame()
  finally:
    if os.path.exists(tmp_path) and tmp_path != latest_file:
      try:
        os.remove(tmp_path)
      except Exception:
        pass


def standardize_slx_df(df):
  df.columns = [str(c).strip().upper() for c in df.columns]
  col_map = {}
  for c in df.columns:
    uc = c.upper()
    if "DATE" in uc or "TIME" in uc or "DAY" in uc:
      col_map["DATE"] = c
    elif (
        "EQUIPMENT" in uc
        or "MACHINE" in uc
        or "ASSET" in uc
        or "RESOURCE" in uc
    ):
      col_map["MACHINE"] = c
    elif "REASON" in uc or "LOSS" in uc or "CAUSE" in uc or "DESCRIPTION" in uc:
      col_map["REASON"] = c
    elif "DURATION" in uc or "DOWNTIME" in uc or "MIN" in uc or "TIME" in uc:
      col_map["DURATION"] = c

  clean_df = pd.DataFrame()
  clean_df["RAW_DATE"] = (
      df[col_map["DATE"]].astype(str).str.strip()
      if "DATE" in col_map
      else df.iloc[:, 0].astype(str).str.strip()
  )
  clean_df["MACHINE"] = (
      df[col_map["MACHINE"]].astype(str).str.strip()
      if "MACHINE" in col_map
      else df.iloc[:, 1].astype(str).str.strip()
  )
  clean_df["REASON"] = (
      df[col_map["REASON"]].astype(str).str.strip()
      if "REASON" in col_map
      else df.iloc[:, 2].astype(str).str.strip()
  )

  dur_col = col_map.get(
      "DURATION", df.columns[3] if len(df.columns) > 3 else df.columns[-1]
  )
  clean_df["DURATION"] = pd.to_numeric(df[dur_col], errors="coerce").fillna(0)
  clean_df["TIMESTAMP"] = pd.to_datetime(clean_df["RAW_DATE"], errors="coerce")
  clean_df["RAW_DATE_STR"] = clean_df["TIMESTAMP"].dt.strftime(
      "%Y-%m-%d"
  ).fillna(clean_df["RAW_DATE"])

  def is_target_making_machine(mach_name):
    m = str(mach_name).upper()
    has_target = any(k in m for k in ["OVEN", "CUTTER", "PDS"])
    is_packing = any(p in m for p in ["PACK", "WRAP", "CARTON", "TTLI", "HCM"])
    return has_target and not is_packing

  return clean_df[clean_df["MACHINE"].apply(is_target_making_machine)].copy()


@st.cache_data
def load_packing_minor_stops(mis_path_or_buffer, ms_dir):
  try:
    files = glob.glob(os.path.join(ms_dir, "**", "*.xls*"), recursive=True)
    valid_files = [
        f for f in files if not os.path.basename(f).startswith("~$")
    ]
    if valid_files:
      latest_file = max(valid_files, key=os.path.getmtime)
      tmp_ext = safe_copy_file(latest_file)
      try:
        raw_ext = pd.read_excel(tmp_ext)
        raw_ext.columns = [str(c).strip().upper() for c in raw_ext.columns]
        return standardize_packing_stops_df(raw_ext)
      finally:
        if os.path.exists(tmp_ext) and tmp_ext != latest_file:
          try:
            os.remove(tmp_ext)
          except Exception:
            pass
  except Exception:
    pass

  try:
    if hasattr(mis_path_or_buffer, "read"):
      xl = pd.ExcelFile(mis_path_or_buffer)
    else:
      if not os.path.exists(mis_path_or_buffer):
        return pd.DataFrame()
      tmp_path = safe_copy_file(mis_path_or_buffer)
      xl = pd.ExcelFile(tmp_path)
      if os.path.exists(tmp_path) and tmp_path != mis_path_or_buffer:
        try:
          os.remove(tmp_path)
        except Exception:
          pass

    ms_sheet = next(
        (
            s
            for s in xl.sheet_names
            if "minor" in s.lower() and "stop" in s.lower()
        ),
        None,
    )
    if ms_sheet:
      raw_ms = pd.read_excel(xl, sheet_name=ms_sheet)
      raw_ms.columns = [str(c).strip().upper() for c in raw_ms.columns]
      return standardize_packing_stops_df(raw_ms)
  except Exception:
    pass

  return pd.DataFrame()


def standardize_packing_stops_df(df):
  df.columns = [str(c).strip().upper() for c in df.columns]
  clean = pd.DataFrame()
  d_col = next((c for c in df.columns if "DATE" in c or "DAY" in c), None)
  m_col = next(
      (
          c
          for c in df.columns
          if "MACHINE" in c or "EQUIPMENT" in c or "ASSET" in c
      ),
      None,
  )
  r_col = next(
      (
          c
          for c in df.columns
          if any(
              k in c
              for k in ["REASON", "FAULT", "DESCRIPTION", "PHENOMENON", "STOP"]
          )
      ),
      None,
  )
  c_col = next(
      (
          c
          for c in df.columns
          if any(
              k in c for k in ["COUNT", "NO OF", "NO. OF", "FREQUENCY", "STOPS"]
          )
      ),
      None,
  )
  l_col = next((c for c in df.columns if "LINE" in c), None)

  if d_col:
    clean["RAW_DATE"] = df[d_col].astype(str).str.strip()
  else:
    clean["RAW_DATE"] = "9-Sep"

  if m_col:
    clean["MACHINE"] = df[m_col].astype(str).str.strip()
  else:
    clean["MACHINE"] = "Perk Packing"

  if r_col:
    clean["REASON"] = df[r_col].astype(str).str.strip()
  else:
    clean["REASON"] = "Minor Stop"

  if c_col:
    clean["COUNT"] = pd.to_numeric(df[c_col], errors="coerce").fillna(1)
  else:
    clean["COUNT"] = 1

  if l_col:
    clean["LINE"] = df[l_col].astype(str).str.strip()
  else:
    clean["LINE"] = "PERK"

  clean["NORM_TAG"] = clean["RAW_DATE"].apply(to_norm_tag)
  clean["TIMESTAMP"] = clean["RAW_DATE"].apply(parse_date_to_timestamp)

  def is_perk_packing(row):
    line_txt = str(row.get("LINE", "")).upper()
    mach_txt = str(row["MACHINE"]).upper()
    is_perk_line = "PERK" in line_txt or "PERK" in mach_txt
    is_pack = any(
        k in mach_txt
        for k in ["HCM", "TTLI", "WRAP", "FLOW", "CARTON", "CASE", "PACK"]
    )
    return is_perk_line and is_pack

  return clean[clean.apply(is_perk_packing, axis=1)]


@st.cache_data
def load_efficiency_data(folder_path_or_buffer):
  if hasattr(folder_path_or_buffer, "read"):
    try:
      df_raw = pd.read_excel(folder_path_or_buffer, header=None)
      header_idx = 0
      for idx, row in df_raw.iterrows():
        if row.dropna().count() >= 3:
          header_idx = idx
          break
      folder_path_or_buffer.seek(0)
      df = pd.read_excel(folder_path_or_buffer, skiprows=header_idx)
      df.columns = [
          str(c).strip() for c in df.columns if pd.notna(c) and str(c) != ""
      ]
      return df.dropna(how="all")
    except Exception:
      return pd.DataFrame()

  if not os.path.exists(folder_path_or_buffer):
    return pd.DataFrame()
  files = glob.glob(os.path.join(folder_path_or_buffer, "**", "*.xls*"), recursive=True)
  valid_files = [f for f in files if not os.path.basename(f).startswith("~$")]
  if not valid_files:
    return pd.DataFrame()
  latest_eff = max(valid_files, key=os.path.getmtime)
  tmp_path = safe_copy_file(latest_eff)
  try:
    df_raw = pd.read_excel(tmp_path, header=None)
    header_idx = 0
    for idx, row in df_raw.iterrows():
      if row.dropna().count() >= 3:
        header_idx = idx
        break
    df = pd.read_excel(tmp_path, skiprows=header_idx)
    df.columns = [
        str(c).strip() for c in df.columns if pd.notna(c) and str(c) != ""
    ]
    return df.dropna(how="all")
  except Exception:
    return pd.DataFrame()
  finally:
    if os.path.exists(tmp_path) and tmp_path != latest_eff:
      try:
        os.remove(tmp_path)
      except Exception:
        pass


try:
  if uploaded_mis is not None:
    mis_source = uploaded_mis
  elif os.path.exists(FILE_PATH):
    mis_source = FILE_PATH
  else:
    mis_source = "PERK MIS.-2026.xlsx"

  rows_data, blocks, all_board_blocks = load_dms_6day_matrix(
      mis_source, SHEET_DMS
  )
except Exception as e:
  st.error(
      "⚠️ Error loading DMS Board source file. Please upload"
      " `PERK MIS.-2026.xlsx` in the sidebar."
  )
  st.stop()

if uploaded_shoplogix is not None:
  try:
    df_shoplogix_all = standardize_slx_df(pd.read_excel(uploaded_shoplogix))
  except Exception:
    df_shoplogix_all = pd.DataFrame()
else:
  df_shoplogix_all = load_shoplogix_from_folder(SHOPLOGIX_DIR)

if uploaded_eff is not None:
  df_eff_master = load_efficiency_data(uploaded_eff)
else:
  df_eff_master = load_efficiency_data(EFFICIENCY_DIR)

if uploaded_minor_stops is not None:
  try:
    df_packing_ms_all = standardize_packing_stops_df(
        pd.read_excel(uploaded_minor_stops)
    )
  except Exception:
    df_packing_ms_all = pd.DataFrame()
else:
  df_packing_ms_all = load_packing_minor_stops(mis_source, MINOR_STOPS_DIR)

monthly_ge_dict = load_ge_loss_monthly(mis_source)

PERCENT_KPIS = [
    "compliance",
    "overweight",
    "waste cost",
    "papas",
    "closure",
    "attendance",
]


def format_cell_value(val_str, kpi_name, uom):
  if not val_str or str(val_str).strip() in ("", "nan", "NAN"):
    return ""
  
  # Prevent error strings or unparsed strings from breaking layout
  val_s = str(val_str).strip()
  if "DIV/0" in val_s.upper():
    return "#DIV/0!"

  cleaned = val_s.replace("%", "").strip()
  is_percent_metric = (
      any(term in kpi_name.lower() for term in PERCENT_KPIS) or (uom == "%")
  )

  if re.match(r"^-?\d+(\.\d+)?$", cleaned):
    num = float(cleaned)
    if (
        is_percent_metric
        and "%" not in val_s
        and num <= 1.0
        and num > 0.0
    ):
      num = num * 100.0
    
    # Weight metrics (like Final Weight, Book Weight, etc.) should keep 1 decimal place cleanly
    if any(w in kpi_name.lower() for w in ["weight", "overweight"]):
      return f"{num:.1f}"

    formatted_num = f"{num:.1f}"
    return f"{formatted_num}%" if is_percent_metric else formatted_num
  return val_s


def get_rag_class(kpi, target, actual):
  if not actual or str(actual).strip() == "" or "#DIV/0!" in str(actual):
    return ""
  act_upper = str(actual).strip().upper()
  if act_upper == "YES":
    return "rag-green"
  elif act_upper == "NO":
    return "rag-red"
  try:
    t_clean = (
        float(str(target).replace("%", "").strip()) if target else None
    )
    a_clean = float(str(actual).replace("%", "").strip())
    if t_clean is None:
      return ""
    lower_is_better = [
        "loss",
        "stop",
        "breakdown",
        "overweight",
        "waste",
        "mttr",
    ]
    if any(term in kpi.lower() for term in lower_is_better):
      return "rag-green" if a_clean <= t_clean else "rag-red"
    else:
      return "rag-green" if a_clean >= t_clean else "rag-red"
  except Exception:
    return ""


def get_pillar_css(p_name):
  p = p_name.lower()
  if "safety" in p:
    return "pillar-safety"
  if "quality" in p:
    return "pillar-quality"
  if "il6s" in p:
    return "pillar-il6s"
  if "delivery" in p:
    return "pillar-delivery"
  if "cost" in p:
    return "pillar-cost"
  if "sustain" in p:
    return "pillar-sustainability"
  return "pillar-morale"


st.markdown('<div class="board-header">PERK DMS BOARD</div>', unsafe_allow_html=True)

col_ctrl1, col_ctrl2 = st.columns([4, 1])
with col_ctrl1:
  available_pillars = ["ALL"] + list(dict.fromkeys(r["pillar"] for r in rows_data))
  selected_pillar = st.selectbox("Filter Pillar View:", available_pillars)

with col_ctrl2:
  st.write("")
  if st.button("🔄 Refresh Data", use_container_width=True):
    st.cache_data.clear()
    st.rerun()

filtered_rows = [
    r
    for r in rows_data
    if (selected_pillar == "ALL" or r["pillar"] == selected_pillar)
]
pillar_counts = {}
for r in filtered_rows:
  p = r["pillar"]
  pillar_counts[p] = pillar_counts.get(p, 0) + 1

html = ['<div class="dms-table-container"><table class="dms-table">']
html.append("<thead><tr>")
html.append('<th rowspan="2" style="width: 130px;">PILLAR</th>')
html.append('<th rowspan="2" style="width: 260px;">KPI / KAI</th>')
html.append('<th rowspan="2" style="width: 80px;">UoM</th>')
for b in blocks:
  html.append(
      f'<th colspan="2">{b["date"]}<br><span style="font-size: 12px;'
      f' font-weight: normal; color: #cbd5e1;">{b["day"]}</span></th>'
  )
html.append("</tr><tr>")
for _ in blocks:
  html.append('<th class="th-sub">TARGET</th>')
  html.append('<th class="th-sub">ACTUAL</th>')
html.append("</tr></thead><tbody>")

rendered_pillars = set()
for r in filtered_rows:
  html.append("<tr>")
  pillar = r["pillar"]
  if pillar not in rendered_pillars:
    span = pillar_counts[pillar]
    pillar_class = get_pillar_css(pillar)
    html.append(
        f'<td class="cell-pillar {pillar_class}" rowspan="{span}">{pillar}</td>'
    )
    rendered_pillars.add(pillar)

  html.append(f'<td class="cell-kpi">{r["kpi"]}</td>')
  html.append(f'<td class="cell-uom">{r["uom"]}</td>')

  for day in r["days"]:
    t_display = format_cell_value(day["target"], r["kpi"], r["uom"])
    a_display = format_cell_value(day["actual"], r["kpi"], r["uom"])

    excel_bg = day.get("excel_bg", "")
    if not excel_bg:
      rag_class = get_rag_class(r["kpi"], t_display, a_display)
      if rag_class == "rag-green":
        excel_bg = "#dcfce7"
        text_color = "#166534"
      elif rag_class == "rag-red":
        excel_bg = "#fee2e2"
        text_color = "#991b1b"
      else:
        text_color = "#1e293b"
    else:
      text_color = "#1e293b"

    actual_style = (
        f'style="background-color: {excel_bg}; color: {text_color}; font-weight:'
        ' 800;"'
        if excel_bg
        else 'style="font-weight: 600;"'
    )

    html.append(f'<td class="cell-target">{t_display}</td>')
    html.append(f"<td {actual_style}>{a_display}</td>")

  html.append("</tr>")

html.append("</tbody></table></div>")
st.markdown("".join(html), unsafe_allow_html=True)

cutter_ge = 0.0
hll_cutter_ge = 0.0
samridhi_packing_ge = 0.0
hll_packing_ge = 0.0
card_date_str = TARGET_DATE.strftime("%d-%b-%Y")

if not df_eff_master.empty:

  def find_col(df, options):
    for opt in options:
      for c in df.columns:
        if opt.lower() == str(c).lower() or opt.lower() in str(c).lower():
          return c
    return df.columns[0]

  c_date_eff = find_col(df_eff_master, ["Date Calendar DateTime", "Date Ca", "Date"])
  c_mach_eff = find_col(df_eff_master, ["Machine Name", "Machine"])
  c_ge_eff = find_col(df_eff_master, ["MDLZ GE%", "MDLZ GE", "GE", "Efficiency"])

  try:
    df_eff_master["Parsed_Date"] = pd.to_datetime(
        df_eff_master[c_date_eff], errors="coerce"
    ).dt.date
    target_dt_obj = TARGET_DATE
    if target_dt_obj in df_eff_master["Parsed_Date"].values:
      active_eff_date = target_dt_obj
    else:
      active_eff_date = df_eff_master["Parsed_Date"].dropna().max()

    card_date_str = pd.to_datetime(active_eff_date).strftime("%d-%b-%Y")
    df_eff_sub = df_eff_master[
        df_eff_master["Parsed_Date"] == active_eff_date
    ]
  except Exception:
    df_eff_sub = df_eff_master

  machine_effs = {}

  for _, row in df_eff_sub.iterrows():
    m_name = str(row[c_mach_eff]).upper()
    raw_val = row[c_ge_eff]
    try:
      ge_val = float(str(raw_val).replace("%", "").strip())
      if ge_val > 1.0:
        ge_val = ge_val / 100.0
    except Exception:
      ge_val = 0.0

    machine_effs[m_name] = ge_val

    if "PERK" in m_name and "HLL" in m_name and "CUTTER" in m_name:
      hll_cutter_ge = ge_val * 100.0
    elif "PERK" in m_name and "CUTTER" in m_name and "HLL" not in m_name:
      cutter_ge = ge_val * 100.0

  samridhi_total_delivered = 0.0
  for i in range(1, 5):
    match_eff = 0.0
    for m_key, m_val in machine_effs.items():
      if f"HCM-{i}" in m_key or f"HCM {i}" in m_key or f"HCM{i}" in m_key:
        match_eff = m_val
        break
    samridhi_total_delivered += match_eff * 600.0
  samridhi_packing_ge = (samridhi_total_delivered / 2400.0) * 100.0

  hll_total_delivered = 0.0
  hll_caps = {1: 400.0, 2: 300.0, 3: 300.0}
  for i, cap in hll_caps.items():
    match_eff = 0.0
    for m_key, m_val in machine_effs.items():
      if f"HLM-{i}" in m_key or f"HLM {i}" in m_key or f"HLM{i}" in m_key:
        match_eff = m_val
        break
    hll_total_delivered += match_eff * cap
  hll_packing_ge = (hll_total_delivered / 1000.0) * 100.0

cutter_bg = "#991b1b" if cutter_ge < 80.0 else "#166534"
hll_cutter_bg = "#991b1b" if hll_cutter_ge < 80.0 else "#166534"
samridhi_bg = "#991b1b" if samridhi_packing_ge < 80.0 else "#166534"
hll_pack_bg = "#991b1b" if hll_packing_ge < 80.0 else "#166534"

st.divider()
main_col, right_col = st.columns([4.2, 1.3])

with main_col:
  chart_col1, chart_col2 = st.columns(2)

  with chart_col1:
    c_card1, c_card2 = st.columns(2)
    with c_card1:
      st.markdown(
          f"""
                <div class="card-container">
                    <div class="card-header">PERK CUTTER GE%</div>
                    <div class="card-body" style="background-color: {cutter_bg}; color: #ffffff;">{cutter_ge:.1f}%</div>
                    <div style="background-color: #f8fafc; color: #475569; font-size: 12px; font-weight: 700; padding: 4px; border-bottom-left-radius: 8px; border-bottom-right-radius: 8px; border-top: 1px solid #e2e8f0;">Target: 80.0%</div>
                </div>
                """,
          unsafe_allow_html=True,
      )
    with c_card2:
      st.markdown(
          f"""
                <div class="card-container">
                    <div class="card-header">PERK HLL CUTTER GE%</div>
                    <div class="card-body" style="background-color: {hll_cutter_bg}; color: #ffffff;">{hll_cutter_ge:.1f}%</div>
                    <div style="background-color: #f8fafc; color: #475569; font-size: 12px; font-weight: 700; padding: 4px; border-bottom-left-radius: 8px; border-bottom-right-radius: 8px; border-top: 1px solid #e2e8f0;">Target: 80.0%</div>
                </div>
                """,
          unsafe_allow_html=True,
      )

    st.markdown(
        '<div class="card-container" style="border: none; box-shadow: none;">',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="card-header" style="font-size: 16px; padding: 10px;'
        ' border-radius: 8px 8px 0 0;">MAKING</div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="chart-card" style="border-top-left-radius: 0;'
        ' border-top-right-radius: 0;">',
        unsafe_allow_html=True,
    )

    if df_shoplogix_all.empty:
      st.info(
          "ℹ️ No Shoplogix making records found. Upload a report in the sidebar"
          " if needed."
      )
    else:
      df_slx = df_shoplogix_all.copy()
      df_slx["NORM_DATE"] = pd.to_datetime(
          df_slx["TIMESTAMP"], errors="coerce"
      ).dt.normalize()

      if not df_slx["NORM_DATE"].dropna().empty:
        target_dt = pd.Timestamp(TARGET_DATE)
        mask = df_slx["NORM_DATE"] == target_dt
        df_last_day = df_slx[mask]

        if df_last_day.empty:
          max_date = df_slx["NORM_DATE"].max()
          df_last_day = df_slx[df_slx["NORM_DATE"] == max_date]
          display_date = str(max_date.date())
        else:
          display_date = str(TARGET_DATE)
      else:
        df_last_day = df_slx
        display_date = "Latest"

      st.caption(f"📅 Target Date: **{display_date}**")

      all_shoplogix_machines = sorted(
          list(set(df_shoplogix_all["MACHINE"].dropna().unique().tolist()))
      )
      selected_slx_machines = st.multiselect(
          "🎯 Filter Making Machines (Multi-select):",
          options=all_shoplogix_machines,
          default=all_shoplogix_machines,
          key="sel_making_mach",
      )

      df_plot_slx = df_last_day.copy()
      if selected_slx_machines and not df_plot_slx.empty:
        df_plot_slx = df_plot_slx[
            df_plot_slx["MACHINE"].isin(selected_slx_machines)
        ]

      chart1_done = False
      if not df_plot_slx.empty:
        df_plot_slx["LABEL"] = (
            df_plot_slx["MACHINE"] + " (" + df_plot_slx["REASON"] + ")"
        )

        top_slx = (
            df_plot_slx.groupby("LABEL")["DURATION"]
            .sum()
            .reset_index()
            .sort_values(by="DURATION", ascending=False)
        )
        top_slx = top_slx[top_slx["DURATION"] > 0].head(5)

        if not top_slx.empty:
          fig1 = go.Figure(
              go.Bar(
                  x=top_slx["LABEL"].astype(str),
                  y=top_slx["DURATION"],
                  text=[f"{v:.1f} m" for v in top_slx["DURATION"]],
                  textposition="auto",
                  marker=dict(
                      color="#DC2626", line=dict(color="#991B1B", width=1.5)
                  ),
              )
          )
          fig1.update_layout(
              template="plotly_white",
              height=320,
              margin=dict(l=20, r=20, t=20, b=50),
              xaxis=dict(tickangle=-20),
              yaxis=dict(title="Downtime (Minutes)"),
          )
          st.plotly_chart(fig1, use_container_width=True)
          chart1_done = True

      if not chart1_done:
        st.info(
            f"ℹ️ Zero Shoplogix making downtimes recorded for `{display_date}`."
        )

    st.markdown("</div></div>", unsafe_allow_html=True)

  with chart_col2:
    c_card3, c_card4 = st.columns(2)
    with c_card3:
      st.markdown(
          f"""
                <div class="card-container">
                    <div class="card-header">SAMRIDHI PERK GE%</div>
                    <div class="card-body" style="background-color: {samridhi_bg}; color: #ffffff;">{samridhi_packing_ge:.1f}%</div>
                    <div style="background-color: #f8fafc; color: #475569; font-size: 12px; font-weight: 700; padding: 4px; border-bottom-left-radius: 8px; border-bottom-right-radius: 8px; border-top: 1px solid #e2e8f0;">Target: 80.0%</div>
                </div>
                """,
          unsafe_allow_html=True,
      )
    with c_card4:
      st.markdown(
          f"""
                <div class="card-container">
                    <div class="card-header">HLL PACKING GE%</div>
                    <div class="card-body" style="background-color: {hll_pack_bg}; color: #ffffff;">{hll_packing_ge:.1f}%</div>
                    <div style="background-color: #f8fafc; color: #475569; font-size: 12px; font-weight: 700; padding: 4px; border-bottom-left-radius: 8px; border-bottom-right-radius: 8px; border-top: 1px solid #e2e8f0;">Target: 80.0%</div>
                </div>
                """,
          unsafe_allow_html=True,
      )

    st.markdown(
        '<div class="card-container" style="border: none; box-shadow: none;">',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="card-header" style="font-size: 16px; padding: 10px;'
        ' border-radius: 8px 8px 0 0;">PACKING</div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="chart-card" style="border-top-left-radius: 0;'
        ' border-top-right-radius: 0;">',
        unsafe_allow_html=True,
    )

    df_packing_day = df_packing_ms_all.copy()
    if not df_packing_day.empty and "NORM_TAG" in df_packing_day.columns:
      target_tag = blocks[0]["norm_tag"] if blocks else "9-SEP"
      if target_tag in df_packing_day["NORM_TAG"].values:
        df_packing_day = df_packing_day[
            df_packing_day["NORM_TAG"] == target_tag
        ]
      elif not df_packing_day["NORM_TAG"].empty:
        df_packing_day.sort_values(by="TIMESTAMP", ascending=True, inplace=True)
        latest_ms_tag = df_packing_day["NORM_TAG"].iloc[-1]
        df_packing_day = df_packing_day[
            df_packing_day["NORM_TAG"] == latest_ms_tag
        ]

    packing_machines = (
        sorted(
            list(set(df_packing_ms_all["MACHINE"].dropna().unique().tolist()))
        )
        if not df_packing_ms_all.empty
        else []
    )
    selected_packing_mach = st.selectbox(
        "🎯 Filter Packing Machine (Perk Line):",
        ["ALL Perk Packing Machines"] + packing_machines,
        key="sel_packing_mach",
    )

    df_plot_packing = df_packing_day.copy()
    if (
        selected_packing_mach != "ALL Perk Packing Machines"
        and not df_plot_packing.empty
    ):
      df_plot_packing = df_plot_packing[
          df_plot_packing["MACHINE"] == selected_packing_mach
      ]

    chart2_done = False
    if not df_plot_packing.empty:
      df_plot_packing["LABEL"] = (
          df_plot_packing["MACHINE"] + " (" + df_plot_packing["REASON"] + ")"
      )

      top_ms = (
          df_plot_packing.groupby("LABEL")["COUNT"]
          .sum()
          .reset_index()
          .sort_values(by="COUNT", ascending=False)
      )
      top_ms = top_ms[top_ms["COUNT"] > 0].head(5)

      if not top_ms.empty:
        fig2 = go.Figure(
            go.Bar(
                x=top_ms["LABEL"].astype(str),
                y=top_ms["COUNT"],
                text=[f"{int(v)}" for v in top_ms["COUNT"]],
                textposition="auto",
                marker=dict(
                    color="#F59E0B", line=dict(color="#B45309", width=1.5)
                ),
            )
        )
        fig2.update_layout(
            template="plotly_white",
            height=320,
            margin=dict(l=20, r=20, t=20, b=50),
            xaxis=dict(tickangle=-20),
            yaxis=dict(title="No. of Minor Stops"),
        )
        st.plotly_chart(fig2, use_container_width=True)
        chart2_done = True

    if not chart2_done:
      st.info("ℹ️ Zero Perk Packing minor stops recorded.")

    st.markdown("</div></div>", unsafe_allow_html=True)

  st.markdown("---")
  st.markdown(
      """
        <div style="background: linear-gradient(135deg, #1f0b35 0%, #3b1764 100%); color: #ffffff; padding: 12px 20px; border-radius: 8px; font-weight: 700; font-size: 1.1rem; margin-bottom: 15px; letter-spacing: 0.8px;">
            📝 MINUTES OF MEETING (MoM) — History & Action Log
        </div>
        """,
      unsafe_allow_html=True,
  )


  def load_mom_history(path, uploaded_file):
    if uploaded_file is not None:
      try:
        return pd.read_excel(uploaded_file)
      except Exception:
        pass
    if os.path.exists(path):
      try:
        return pd.read_excel(path)
      except Exception:
        pass
    return pd.DataFrame([{
        "Date": TARGET_DATE.strftime("%d-%b-%Y"),
        "Agenda": "Cutter downtime & minor stops review",
        "Action Item": "Check chain tensioner and guide rails",
        "Owner": "Dharmendra / Arvind",
        "Target Date": (TARGET_DATE + timedelta(days=2)).strftime("%d-%b-%Y"),
        "Status": "In Progress",
    }])


  if "mom_df" not in st.session_state or uploaded_mom is not None:
    st.session_state.mom_df = load_mom_history(MOM_FILE_PATH, uploaded_mom)

  col_f1, col_f2 = st.columns([2, 2])
  with col_f1:
    status_filter = st.multiselect(
        "Filter MoM by Status:",
        options=["In Progress", "Completed", "Pending"],
        default=["In Progress", "Completed", "Pending"],
        key="mom_status_filter",
    )
  with col_f2:
    owner_search = st.text_input("Search by Owner / Keyword:", "", key="mom_search")

  view_df = st.session_state.mom_df.copy()
  if status_filter:
    view_df = view_df[view_df["Status"].isin(status_filter)]
  if owner_search:
    mask = view_df.astype(str).apply(
        lambda row: row.str.contains(owner_search, case=False).any(), axis=1
    )
    view_df = view_df[mask]

  st.caption(
      f"Showing {len(view_df)} of {len(st.session_state.mom_df)} total records"
  )

  edited_mom_df = st.data_editor(
      view_df, num_rows="dynamic", use_container_width=True, key="mom_table_editor"
  )

  col_save1, col_save2 = st.columns([1, 4])
  with col_save1:
    if st.button("💾 Save MoM", use_container_width=True):
      st.session_state.mom_df = edited_mom_df
      try:
        edited_mom_df.to_excel(MOM_FILE_PATH, index=False)
        upload_file_to_github(MOM_FILE_PATH, "PERK_MOM_History.xlsx")
        st.success("✅ MoM changes saved locally and synced to GitHub!")
      except Exception as ex:
        st.error(f"⚠️ Failed to save file: {ex}")
  with col_save2:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
      st.session_state.mom_df.to_excel(writer, index=False)
    excel_data = output.getvalue()

    st.download_button(
        label="📥 Download Updated MoM Excel File",
        data=excel_data,
        file_name="PERK_MOM_History.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        use_container_width=True,
    )

with right_col:
  st.markdown(
      """
        <div style="background-color: #250e3e; color: #ffffff; text-align: center; padding: 12px; border-radius: 8px; font-weight: 700; font-size: 13px; margin-bottom: 14px; letter-spacing: 0.8px;">
            MONTHLY GE% (2026)
        </div>
        """,
      unsafe_allow_html=True,
  )

  months_display = [
      "JAN",
      "FEB",
      "MAR",
      "APR",
      "MAY",
      "JUN",
      "JUL",
      "AUG",
      "SEP",
      "OCT",
      "NOV",
      "DEC",
  ]
  for m in months_display:
    val = monthly_ge_dict.get(m, "")
    if val != "" and val is not None:
      try:
        num_val = float(val)
        bg_col = "#166534" if num_val >= 80.0 else "#991b1b"
        display_txt = f"{num_val:.1f}%"
      except Exception:
        bg_col = "#475569"
        display_txt = "N/A"
    else:
      bg_col = "#cbd5e1"
      display_txt = "—"

    text_color = "#1e293b" if bg_col == "#cbd5e1" else "#ffffff"

    st.markdown(
        f"""
            <div class="card-container" style="margin-bottom: 8px;">
                <div class="card-header" style="padding: 4px; font-size: 11px;">{m} GE% (Target: 80%)</div>
                <div class="card-body" style="padding: 6px; font-size: 17px; background-color: {bg_col}; color: {text_color};">{display_txt}</div>
            </div>
            """,
        unsafe_allow_html=True,
    )
